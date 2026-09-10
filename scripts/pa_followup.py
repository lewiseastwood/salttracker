#!/usr/bin/env python3
"""Pennsylvania follow-up: contacts workbook and Right-to-Know drafts.

Pilot set is the five counties with the largest FY2027 contracted tons on the
statewide COSTARS award. Those awards are already in the tracker. The letters
ask each county for *its own* sodium-chloride supply contract (or a statement
that it buys only through COSTARS), which is the gap the statewide scrape cannot
close.

Default is draft-only. This script does not send mail unless you pass --send
and set SALTTRACKER_FOLLOWUP_CONFIRM=YES. Replies can be polled from IMAP and
forwarded to SALTTRACKER_ALERT_EMAIL; that also requires explicit flags.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import imaplib
import os
import smtplib
import sys
from email.message import EmailMessage
from email.parser import BytesParser
from email.policy import default as email_policy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
CONTACTS = ROOT / "data" / "pa_followup" / "contacts.csv"
OUT = ROOT / "data" / "output" / "pa_followup"
DRAFTS = OUT / "drafts"

SUBJECT = "Right-to-Know request: road salt / sodium chloride supply contract"

BODY = """{greeting}

I am requesting public records under Pennsylvania's Right-to-Know Law (65 P.S. § 67.101 et seq.) for academic research on contracted road-salt prices.

Please provide, for the current or most recently awarded winter season:

1. The county's (or its participating municipalities') sodium chloride / road-salt supply contract or purchase order, including awarded vendor, unit price, and contracted tons.
2. If the county buys only through the Commonwealth COSTARS sodium chloride contract and holds no separate award, a short written confirmation of that.

I am not seeking bid bonds, sealed proposals that remain unopened, or any record that is not public. Electronic copies (PDF) are preferred.

Please send records or a response to: {reply_to}

Thank you for your time.

{sender_name}
{sender_org}
"""


def load_contacts(path: Path = CONTACTS) -> list[dict]:
    with path.open(newline="") as fh:
        return list(csv.DictReader(fh))


def greeting_line(row: dict) -> str:
    officer = (row.get("aoro_officer") or "").strip()
    role = (row.get("aoro_title") or "Agency Open Records Officer").strip()
    county = row.get("county") or "the county"
    if officer:
        return f"Dear {officer}, {role}, {county} County:"
    return f"Dear {role}, {county} County:"


def sender_fields() -> tuple[str, str, str]:
    reply = os.environ.get("SALTTRACKER_FOLLOWUP_REPLY_TO", "").strip()
    name = os.environ.get("SALTTRACKER_FOLLOWUP_NAME", "").strip() or "Lewis Eastwood"
    org = os.environ.get(
        "SALTTRACKER_FOLLOWUP_ORG",
        "Academic research — Road salt contract tracker",
    ).strip()
    return reply, name, org


def aoro_email(row: dict) -> str | None:
    status = (row.get("aoro_status") or "").strip().upper()
    email = (row.get("aoro_email") or "").strip()
    if status != "VERIFIED":
        return None
    if not email or email.upper() == "UNVERIFIED" or "@" not in email:
        return None
    return email


def draft_message(row: dict, stamp: str) -> EmailMessage:
    reply, name, org = sender_fields()
    if not reply:
        reply = os.environ.get("SALTTRACKER_ALERT_EMAIL", "").strip() or "SET SALTTRACKER_FOLLOWUP_REPLY_TO"
    msg = EmailMessage()
    msg["Subject"] = SUBJECT
    msg["From"] = os.environ.get("SALTTRACKER_SMTP_FROM", reply)
    dest = aoro_email(row)
    if dest:
        msg["To"] = dest
    else:
        msg["To"] = "UNVERIFIED-DO-NOT-SEND"
    msg["Date"] = stamp
    msg["X-SaltTracker-County"] = row["county"]
    msg["X-SaltTracker-AORO-Status"] = row.get("aoro_status") or ""
    msg.set_content(BODY.format(
        greeting=greeting_line(row),
        reply_to=reply,
        sender_name=name,
        sender_org=org,
    ))
    return msg


def write_workbook(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PA procurement contacts"
    header = list(rows[0].keys()) if rows else []
    fill = PatternFill("solid", fgColor="1B3A4B")
    font = Font(color="FFFFFF", bold=True)
    ws.append(header)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
    for row in rows:
        ws.append([row.get(c, "") for c in header])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(48, max(14, len(str(col[0].value or "")) + 4))
    note = wb.create_sheet("How to use")
    note["A1"] = (
        "Pilot: five Pennsylvania counties by FY2027 contracted tons. "
        "Primary recipient is the county Agency Open Records Officer (RTKL), "
        "verified 10 Sep 2026 from the county Right-to-Know page. Purchasing "
        "contacts are secondary only. Washington AORO email is UNVERIFIED "
        "(not published; use the county web form). Default: python scripts/pa_followup.py "
        "writes this workbook and .eml drafts. Sending requires --send and "
        "SALTTRACKER_FOLLOWUP_CONFIRM=YES and skips UNVERIFIED rows."
    )
    note["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    note.row_dimensions[1].height = 80
    note.column_dimensions["A"].width = 88
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_drafts(rows: list[dict], stamp: str) -> list[Path]:
    DRAFTS.mkdir(parents=True, exist_ok=True)
    written = []
    for row in rows:
        msg = draft_message(row, stamp)
        slug = row["county"].lower().replace(" ", "_")
        path = DRAFTS / f"{stamp[:10]}_{slug}.eml"
        path.write_bytes(bytes(msg))
        written.append(path)
    return written


def send_drafts(rows: list[dict], stamp: str) -> list[str]:
    if os.environ.get("SALTTRACKER_FOLLOWUP_CONFIRM", "").strip() != "YES":
        raise RuntimeError("Refusing to send: set SALTTRACKER_FOLLOWUP_CONFIRM=YES")
    host = os.environ.get("SALTTRACKER_SMTP_HOST", "").strip()
    if not host:
        raise RuntimeError("SALTTRACKER_SMTP_HOST is required to send")
    port = int(os.environ.get("SALTTRACKER_SMTP_PORT", "587"))
    user = os.environ.get("SALTTRACKER_SMTP_USER", "")
    password = os.environ.get("SALTTRACKER_SMTP_PASS", "")
    notes = []
    with smtplib.SMTP(host, port, timeout=20) as smtp:
        smtp.starttls()
        if user:
            smtp.login(user, password)
        for row in rows:
            if not aoro_email(row):
                notes.append(f"skipped {row['county']}: AORO email UNVERIFIED")
                continue
            msg = draft_message(row, stamp)
            smtp.send_message(msg)
            notes.append(f"sent {row['county']} -> {aoro_email(row)}")
    return notes


def poll_inbox() -> list[str]:
    """Forward unseen replies that look like follow-up responses."""
    host = os.environ.get("SALTTRACKER_IMAP_HOST", "").strip()
    user = os.environ.get("SALTTRACKER_IMAP_USER", "").strip()
    password = os.environ.get("SALTTRACKER_IMAP_PASS", "").strip()
    dest = os.environ.get("SALTTRACKER_ALERT_EMAIL", "").strip()
    smtp_host = os.environ.get("SALTTRACKER_SMTP_HOST", "").strip()
    if not (host and user and password and dest and smtp_host):
        raise RuntimeError("IMAP poll needs IMAP_* credentials, ALERT_EMAIL, and SMTP_HOST")
    mailbox = os.environ.get("SALTTRACKER_IMAP_MAILBOX", "INBOX")
    notes = []
    with imaplib.IMAP4_SSL(host) as imap:
        imap.login(user, password)
        imap.select(mailbox)
        typ, data = imap.search(None, "UNSEEN")
        if typ != "OK":
            return ["imap search failed"]
        ids = data[0].split()
        for msg_id in ids:
            typ, payload = imap.fetch(msg_id, "(RFC822)")
            if typ != "OK" or not payload or not payload[0]:
                continue
            raw = payload[0][1]
            parsed = BytesParser(policy=email_policy).parsebytes(raw)
            subject = parsed.get("Subject", "")
            if "road salt" not in subject.lower() and "right-to-know" not in subject.lower():
                continue
            fwd = EmailMessage()
            fwd["Subject"] = f"Fwd (PA follow-up): {subject}"
            fwd["From"] = os.environ.get("SALTTRACKER_SMTP_FROM", dest)
            fwd["To"] = dest
            fwd.set_content(
                f"Forwarded county reply.\n\nFrom: {parsed.get('From')}\n"
                f"Subject: {subject}\n\n{parsed.get_content()}"
            )
            port = int(os.environ.get("SALTTRACKER_SMTP_PORT", "587"))
            with smtplib.SMTP(smtp_host, port, timeout=20) as smtp:
                smtp.starttls()
                smtp_user = os.environ.get("SALTTRACKER_SMTP_USER", "")
                smtp_pass = os.environ.get("SALTTRACKER_SMTP_PASS", "")
                if smtp_user:
                    smtp.login(smtp_user, smtp_pass)
                smtp.send_message(fwd)
            notes.append(f"forwarded {subject}")
    return notes


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--send", action="store_true",
                        help="Send drafts via SMTP (also requires SALTTRACKER_FOLLOWUP_CONFIRM=YES)")
    parser.add_argument("--poll-inbox", action="store_true",
                        help="Forward unseen IMAP replies to SALTTRACKER_ALERT_EMAIL")
    args = parser.parse_args(argv)

    rows = load_contacts()
    if not rows:
        print("No contacts in", CONTACTS, file=sys.stderr)
        return 1
    stamp = dt.datetime.now().replace(microsecond=0).isoformat()
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx = OUT / "PA_procurement_contacts.xlsx"
    write_workbook(rows, xlsx)
    drafts = write_drafts(rows, stamp)
    print(f"Wrote {xlsx}")
    print(f"Wrote {len(drafts)} drafts under {DRAFTS}")
    for p in drafts:
        print(" ", p.name)

    if args.send:
        for line in send_drafts(rows, stamp):
            print(line)
    else:
        print("No mail sent (draft only). Pass --send only after reviewing the .eml files.")

    if args.poll_inbox:
        for line in poll_inbox():
            print(line)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
